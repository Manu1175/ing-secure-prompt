from __future__ import annotations

from pathlib import Path

import openpyxl
import yaml

from secureprompt.prompt_mining import baseline


def test_build_baseline_counts(tmp_path: Path) -> None:
    # Placeholder catalog
    placeholders = [
        {
            "label": "EMAIL",
            "c_level": "C3",
            "templates": ["<CUSTOMER_EMAIL>"]
        },
        {
            "label": "PHONE",
            "c_level": "C4",
            "templates": ["<CUSTOMER_PHONE>"]
        },
    ]
    placeholders_path = tmp_path / "placeholders.yml"
    placeholders_path.write_text(yaml.safe_dump(placeholders), encoding="utf-8")

    prompts_dir = tmp_path / "PROMPTS"
    prompts_dir.mkdir()

    workbook_path = prompts_dir / "example.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Original Prompt", "Sanitized Prompt", "Response", "Sanitized Response"])
    sheet.append([
        "Email john@bank.com",
        "Email <CUSTOMER_EMAIL>",
        "Call on +321234",
        "Call <CUSTOMER_PHONE>",
    ])
    workbook.save(workbook_path)

    data = baseline.build_baseline(str(prompts_dir), placeholders_path=placeholders_path)

    assert data["total_files"] == 1
    file_stats = data["by_file"]["example.xlsx"]
    assert file_stats["by_label"]["EMAIL"] == 1
    assert file_stats["by_label"]["PHONE"] == 1
    assert file_stats["by_c_level"]["C3"] == 1
    assert file_stats["by_c_level"]["C4"] == 1

    rows = baseline.flatten_counts(data)
    assert {row["label"] for row in rows} == {"EMAIL", "PHONE"}
