from __future__ import annotations

import importlib
import json
from pathlib import Path

import openpyxl
import pytest
import yaml

pytest.importorskip("fastapi")

from fastapi.testclient import TestClient


@pytest.fixture()
def ui_client_with_baseline(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    baseline_path = tmp_path / "baseline.json"
    baseline_path.write_text(
        json.dumps(
            {
                "total_files": 1,
                "by_file": {
                    "example.xlsx": {
                        "by_label": {"EMAIL": 1},
                        "by_c_level": {"C3": 1},
                    }
                },
                "global": {"by_label": {"EMAIL": 1}, "by_c_level": {"C3": 1}},
            }
        ),
        encoding="utf-8",
    )

    placeholders_path = tmp_path / "placeholders.yml"
    yaml.safe_dump(
        [{"label": "EMAIL", "c_level": "C3", "templates": ["<CUSTOMER_EMAIL>"]}],
        placeholders_path.open("w", encoding="utf-8"),
    )

    monkeypatch.setenv("SECUREPROMPT_BASELINE_PATH", str(baseline_path))
    monkeypatch.setenv("SECUREPROMPT_PLACEHOLDERS_PATH", str(placeholders_path))

    from secureprompt.ui import scoring

    importlib.reload(scoring)
    scoring.reset_cache()

    import api.main as api_main

    importlib.reload(api_main)
    return TestClient(api_main.app)


def test_ui_scoring_card(ui_client_with_baseline: TestClient, tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Original Prompt", "Sanitized Prompt"])
    sheet.append(["Email a@b.com", "Email <CUSTOMER_EMAIL>"])

    path = tmp_path / "example.xlsx"
    workbook.save(path)
    with path.open("rb") as handle:
        files = {
            "upload": (
                "example.xlsx",
                handle.read(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

    response = ui_client_with_baseline.post(
        "/ui/scrub", data={"clearance": "C1"}, files=files
    )
    assert response.status_code == 200
    body = response.text
    assert "Expected vs Achieved" in body
    assert "Download redacted" in body
    assert "Original SHA256" in body
