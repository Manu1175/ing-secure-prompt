"""Baseline placeholder statistics mined from PROMPTS workbooks."""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, MutableMapping, Tuple

import openpyxl
import yaml

PLACEHOLDER_PATTERN = re.compile(r"<[^>]+>")


def load_placeholder_catalog(placeholders_path: Path) -> Tuple[Dict[str, str], Dict[str, str]]:
    """Return lookup tables for placeholder templates and labels."""

    if not placeholders_path.exists():
        return {}, {}

    with placeholders_path.open("r", encoding="utf-8") as handle:
        items = yaml.safe_load(handle) or []

    template_to_label: Dict[str, str] = {}
    label_to_level: Dict[str, str] = {}
    for entry in items:
        if not isinstance(entry, dict):
            continue
        label = entry.get("label")
        c_level = entry.get("c_level") or entry.get("clearance") or "C3"
        templates = entry.get("templates") or []
        if label:
            label_to_level[label] = str(c_level).upper()
            for template in templates:
                template_to_label[str(template)] = label

    return template_to_label, label_to_level


def _iter_workbook_prompts(path: Path) -> Iterable[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(h).strip() if h is not None else "" for h in next(rows)]
            except StopIteration:
                continue

            header_map = {name.lower(): idx for idx, name in enumerate(headers)}

            prompt_cols = []
            response_cols = []
            for name, idx in header_map.items():
                if "sanitized prompt" in name:
                    prompt_cols.append(idx)
                if "sanitized response" in name:
                    response_cols.append(idx)

            for row in rows:
                if prompt_cols:
                    for idx in prompt_cols:
                        value = row[idx] if idx < len(row) else None
                        text = str(value).strip() if value not in (None, "") else ""
                        if text:
                            yield text
                if response_cols:
                    for idx in response_cols:
                        value = row[idx] if idx < len(row) else None
                        text = str(value).strip() if value not in (None, "") else ""
                        if text:
                            yield text
    finally:
        workbook.close()


def _extract_placeholders(text: str) -> List[str]:
    return PLACEHOLDER_PATTERN.findall(text)


def build_baseline(
    prompts_folder: str,
    *,
    placeholders_path: Path = Path("policy/placeholders.yml"),
) -> Dict[str, Dict[str, Dict[str, int]]]:
    """Compute placeholder label and clearance distributions from sanitized prompts."""

    prompts_dir = Path(prompts_folder)
    if not prompts_dir.is_dir():
        raise FileNotFoundError(f"Prompts directory not found: {prompts_dir}")

    template_to_label, label_to_level = load_placeholder_catalog(placeholders_path)

    by_file: Dict[str, Dict[str, MutableMapping[str, int]]] = {}
    global_by_label: Counter[str] = Counter()
    global_by_level: Counter[str] = Counter()

    workbooks = sorted(prompts_dir.glob("*.xlsx"))
    for workbook_path in workbooks:
        label_counts: Counter[str] = Counter()
        level_counts: Counter[str] = Counter()

        for text in _iter_workbook_prompts(workbook_path):
            for placeholder in _extract_placeholders(text):
                label = template_to_label.get(placeholder)
                if not label:
                    key = placeholder.strip("<>").upper()
                    label = key
                level = label_to_level.get(label, "C3")
                label_counts[label] += 1
                level_counts[level] += 1

        if label_counts:
            by_file[workbook_path.name] = {
                "by_label": dict(label_counts),
                "by_c_level": dict(level_counts),
            }
            global_by_label.update(label_counts)
            global_by_level.update(level_counts)

    baseline = {
        "total_files": len(workbooks),
        "by_file": by_file,
        "global": {
            "by_label": dict(global_by_label),
            "by_c_level": dict(global_by_level),
        },
    }
    return baseline


def write_json(path: Path, data: Dict[str, Any]) -> None:
    """Persist baseline statistics to JSON."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(data, handle, indent=2, sort_keys=True)


def write_csv(path: Path, rows: Iterable[Dict[str, Any]]) -> None:
    """Persist flattened baseline rows to CSV."""
    import csv

    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if not rows:
        path.write_text("", encoding="utf-8")
        return

    fieldnames = list(rows[0].keys())
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def flatten_counts(baseline: Dict[str, Any]) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for filename, data in baseline.get("by_file", {}).items():
        for label, count in data.get("by_label", {}).items():
            records.append({"filename": filename, "label": label, "count": count})
    return records


__all__ = [
    "build_baseline",
    "write_json",
    "write_csv",
    "flatten_counts",
    "load_placeholder_catalog",
]
